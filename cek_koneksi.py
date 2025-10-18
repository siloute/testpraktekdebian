import tkinter as tk
from tkinter import ttk, scrolledtext, filedialog, messagebox, simpledialog
import subprocess
import threading
import time
import platform
import socket
import urllib.request
from urllib.error import URLError
import dns.resolver
from collections import OrderedDict
import openpyxl

class ClientStatusBox(tk.LabelFrame):
    """
    Kotak visual untuk memonitor satu client, kini dengan penilaian otomatis.
    """
    def __init__(self, parent, ip, domain, name, log_callback):
        self.name = name
        self.ip = ip
        
        title_widget = ttk.Frame(parent)
        if name:
            ttk.Label(title_widget, text=name, font=("Helvetica", 10, "bold")).pack(anchor="w")
            ttk.Label(title_widget, text=ip, font=("Helvetica", 8)).pack(anchor="w")
        else:
            ttk.Label(title_widget, text=ip, font=("Helvetica", 10, "bold")).pack(anchor="w")

        super().__init__(parent, labelwidget=title_widget, padx=10, pady=10)
        
        self.log_callback = log_callback
        self.is_checking = False
        self.check_thread = None
        self.previous_status = "IDLE"
        self.score = 0
        self.results = {}

        # Peta Poin Penilaian
        self.score_map = {
            'ping_ip': 5, 'http_ip': 10, 'apache_default': 5,
            'dns_server': 10, 'domain_resolve': 25, 'http_domain': 15
        }

        # --- Elemen UI ---
        self.indicator_panel = tk.Frame(self, height=40, bg="#cccccc")
        self.indicator_panel.pack(fill="x", pady=(0, 10))

        # --- Frame Atas (Domain & Nilai) ---
        top_frame = ttk.Frame(self)
        top_frame.pack(fill="x", pady=(0, 8))

        ttk.Label(top_frame, text="Domain:", font=("Helvetica", 8)).pack(side="left")
        self.domain_entry = ttk.Entry(top_frame, font=("Helvetica", 9))
        self.domain_entry.pack(side="left", fill="x", expand=True, padx=(5, 10))
        if domain: self.domain_entry.insert(0, domain)
        
        self.score_label = ttk.Label(top_frame, text="Nilai: 0", font=("Helvetica", 10, "bold"))
        self.score_label.pack(side="right")

        # --- Frame Hasil ---
        results_frame = ttk.Frame(self, padding=(5, 5))
        results_frame.pack(fill="x", pady=(5, 10))
        self.check_items = OrderedDict([
            ('ping_ip', 'Ping IP'), ('http_ip', 'HTTP (IP)'), ('apache_default', 'Halaman Default'),
            ('dns_server', 'DNS (53)'), ('domain_resolve', 'Resolve'), ('http_domain', 'HTTP (Domain)')
        ])
        self.result_labels = {}
        for i, (key, text) in enumerate(self.check_items.items()):
            ttk.Label(results_frame, text=f"{text}:", font=("Helvetica", 8)).grid(row=i, column=0, sticky="w", pady=1)
            status_label = ttk.Label(results_frame, text="...", font=("Helvetica", 8, "bold"), anchor="e")
            status_label.grid(row=i, column=1, sticky="e", pady=1)
            self.result_labels[key] = status_label
            results_frame.grid_columnconfigure(1, weight=1)

        # --- Tombol ---
        button_frame = ttk.Frame(self)
        button_frame.pack(fill="x")
        self.start_button = ttk.Button(button_frame, text="Start", command=self.start_check, style='small.TButton')
        self.start_button.pack(side="left", fill="x", expand=True, padx=(0, 5))
        self.stop_button = ttk.Button(button_frame, text="Stop", command=self.stop_check, state="disabled", style='small.TButton')
        self.stop_button.pack(side="right", fill="x", expand=True, padx=(5, 0))
        
        ttk.Style().configure('small.TButton', font=('Helvetica', 8))

    def start_check(self):
        if self.is_checking: return
        self.is_checking = True
        self.start_button.config(state="disabled")
        self.stop_button.config(state="normal")
        self.indicator_panel.config(bg="#3498db")
        self.check_thread = threading.Thread(target=self._check_loop, daemon=True)
        self.check_thread.start()

    def stop_check(self):
        if not self.is_checking: return
        self.is_checking = False
        self.start_button.config(state="normal")
        self.stop_button.config(state="disabled")
        if self.previous_status == "SUCCESS": self.log_callback(f"TERPUTUS (Manual): {self.ip}")
        self.update_ui("IDLE", {})
        self.previous_status = "IDLE"

    def update_ui(self, status_code, results):
        self.results = results
        color_map = {"SUCCESS": "#2ecc71", "PARTIAL": "#f39c12", "FAILED": "#e74c3c", "IDLE": "#cccccc"}
        self.indicator_panel.config(bg=color_map.get(status_code, "#cccccc"))
        
        score = 30 if status_code != "IDLE" else 0
        for key, value in results.items():
            if value is True: score += self.score_map.get(key, 0)
        self.score = score
        self.score_label.config(text=f"Nilai: {self.score}")
        
        for key, label in self.result_labels.items():
            res = results.get(key)
            if res is True: text, color = "✅ OK", "green"
            elif res is False: text, color = "❌ GAGAL", "red"
            elif res == 'N/A': text, color = "N/A", "grey"
            else: text, color = "...", "black"
            label.config(text=text, foreground=color)

        if status_code != self.previous_status:
            log_msg = ""
            if status_code == "SUCCESS": log_msg = f"TERKONEKSI: {self.ip}"
            elif self.previous_status == "SUCCESS": log_msg = f"TERPUTUS: {self.ip}"
            if log_msg: self.log_callback(log_msg)
            self.previous_status = status_code

    def _check_loop(self):
        while self.is_checking:
            overall_status, results = self._perform_all_checks()
            self.after(0, self.update_ui, overall_status, results)
            time.sleep(5)
    
    def _perform_all_checks(self):
        domain = self.domain_entry.get().strip()
        results = OrderedDict.fromkeys(self.check_items.keys())
        
        results['ping_ip'] = self._check_ping(self.ip)
        if not results['ping_ip']:
            results['ping_ip'] = False
            return "FAILED", results
        
        results['http_ip'] = self._check_port(self.ip, 80)
        results['dns_server'] = self._check_port(self.ip, 53)
        results['apache_default'] = self._check_apache_default(self.ip, domain) if results['http_ip'] else False
        
        if domain:
            results['domain_resolve'] = self._check_domain_resolution(domain, self.ip)
            results['http_domain'] = self._check_port(domain, 80) if results['domain_resolve'] else False
        else:
            results['domain_resolve'] = 'N/A'
            results['http_domain'] = 'N/A'
            
        is_success = all(v for v in results.values() if v is True)
        if is_success: return "SUCCESS", results
        return "PARTIAL", results

    def _check_ping(self, host):
        param = '-n' if platform.system().lower() == 'windows' else '-c'
        command = ['ping', param, '1', '-w', '1000', host]
        try:
            startupinfo = subprocess.STARTUPINFO() if platform.system().lower() == 'windows' else None
            if startupinfo:
                startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
            return subprocess.run(command, capture_output=True, text=True, startupinfo=startupinfo, timeout=2).returncode == 0
        except (subprocess.TimeoutExpired, FileNotFoundError):
            return False

    def _check_port(self, host, port):
        try:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
                sock.settimeout(2)
                return sock.connect_ex((host, port)) == 0
        except (socket.gaierror, OSError):
            return False

    def _check_apache_default(self, ip, domain):
        try:
            url = f"http://{ip}"
            headers = {'Host': domain} if domain else {}
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=3) as response:
                html = response.read().decode('utf-8', errors='ignore')
                return "<title>Apache2 Debian Default Page</title>" not in html
        except (URLError, socket.timeout):
            return False

    def _check_domain_resolution(self, domain, expected_ip):
        try:
            resolver = dns.resolver.Resolver()
            resolver.nameservers = [expected_ip]
            resolver.timeout = 2
            resolver.lifetime = 2
            answers = resolver.resolve(domain, 'A')
            return any(rdata.address == expected_ip for rdata in answers)
        except Exception:
            return False

class ManualInputDialog(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("Tambah Siswa Manual")
        self.transient(parent)
        self.grab_set()
        self.result = None

        f = ttk.Frame(self, padding="15")
        f.pack(expand=True, fill="both")

        ttk.Label(f, text="Alamat IP (Wajib):").grid(row=0, column=0, sticky="w", pady=4)
        self.ip_entry = ttk.Entry(f, width=35)
        self.ip_entry.grid(row=0, column=1, sticky="ew", pady=4)
        self.ip_entry.focus_set()

        ttk.Label(f, text="Nama Siswa:").grid(row=1, column=0, sticky="w", pady=4)
        self.name_entry = ttk.Entry(f, width=35)
        self.name_entry.grid(row=1, column=1, sticky="ew", pady=4)

        ttk.Label(f, text="Domain Web:").grid(row=2, column=0, sticky="w", pady=4)
        self.domain_entry = ttk.Entry(f, width=35)
        self.domain_entry.grid(row=2, column=1, sticky="ew", pady=4)

        f.grid_columnconfigure(1, weight=1)

        bf = ttk.Frame(self, padding=(0, 10, 15, 10))
        bf.pack(fill="x")
        
        ok_button = ttk.Button(bf, text="OK", command=self.apply, style="Accent.TButton")
        ok_button.pack(side="right")
        ttk.Button(bf, text="Batal", command=self.destroy).pack(side="right", padx=10)
        
        self.bind("<Return>", lambda e: self.apply())
        self.protocol("WM_DELETE_WINDOW", self.destroy)

    def apply(self):
        ip = self.ip_entry.get().strip()
        if not ip:
            messagebox.showwarning("Input Tidak Lengkap", "Alamat IP wajib diisi.", parent=self)
            return
        self.result = {"ip": ip, "name": self.name_entry.get().strip(), "domain": self.domain_entry.get().strip()}
        self.destroy()

class StatusMonitorApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Aplikasi Monitor Status Client (v13 - Scoring & Export)")
        self.geometry("1400x800")
        
        self.status_boxes = []
        self.timer_job_id = None

        pw = ttk.PanedWindow(self, orient=tk.VERTICAL)
        pw.pack(fill="both", expand=True)
        
        top_frame = ttk.Frame(pw, padding=10)
        pw.add(top_frame, weight=4)
        
        self._create_control_panel(top_frame)
        self._create_scrollable_area(top_frame)

        log_frame = ttk.LabelFrame(pw, text="Log Aktivitas", padding=10)
        pw.add(log_frame, weight=1)
        self.log_area = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD, state='disabled', font=("Consolas", 10))
        self.log_area.pack(fill="both", expand=True)

    def _create_control_panel(self, parent):
        cf = ttk.LabelFrame(parent, text="Kontrol", padding="10")
        cf.pack(fill="x", pady=(0, 10))

        ig = ttk.Frame(cf); ig.pack(side="left", padx=10)
        ttk.Button(ig, text="Import Excel", command=self.import_from_excel, style="Accent.TButton").pack(side="left", padx=5, ipady=4)
        ttk.Button(ig, text="Tambah Manual", command=self.add_manual_entry).pack(side="left", padx=5)

        ttk.Separator(cf, orient='vertical').pack(side="left", fill='y', padx=10)

        ag = ttk.Frame(cf); ag.pack(side="left", padx=10)
        self.start_all_button = ttk.Button(ag, text="Start Semua", command=self.start_all, state="disabled")
        self.start_all_button.pack(side="left", padx=5)
        self.stop_all_button = ttk.Button(ag, text="Stop Semua", command=self.stop_all, state="disabled")
        self.stop_all_button.pack(side="left", padx=5)

        ttk.Separator(cf, orient='vertical').pack(side="left", fill='y', padx=10)

        ttk.Button(cf, text="Export Hasil", command=self.export_results, style="Accent.TButton").pack(side="left", ipady=4, padx=5)

        sg = ttk.Frame(cf); sg.pack(side="left", padx=10)
        ttk.Label(sg, text="Kolom:").pack(side="left")
        self.columns_var = tk.StringVar(value="5")
        ttk.Spinbox(sg, from_=1, to=12, textvariable=self.columns_var, width=5, command=self.rearrange_grid).pack(side="left")

        tg = ttk.Frame(cf); tg.pack(side="right", padx=10)
        ttk.Label(tg, text="Timer:").pack(side="left")
        self.timer_var = tk.StringVar(value="10")
        self.timer_spinbox = ttk.Spinbox(tg, from_=1, to=180, textvariable=self.timer_var, width=5)
        self.timer_spinbox.pack(side="left", padx=5)
        self.timer_button = ttk.Button(tg, text="Start Timer", command=self.toggle_countdown)
        self.timer_button.pack(side="left")
        self.timer_display_label = ttk.Label(tg, text="00:00", font=("Consolas", 12, "bold"), foreground="blue")
        self.timer_display_label.pack(side="left", padx=10)
        
        ttk.Style().configure('Accent.TButton', font=('Helvetica', 10, 'bold'))

    def _create_scrollable_area(self, parent):
        cf = ttk.Frame(parent)
        cf.pack(fill="both", expand=True)
        
        # --- Create Canvas with both scrollbars ---
        self.canvas = tk.Canvas(cf, highlightthickness=0)
        
        self.scrollbar_v = ttk.Scrollbar(cf, orient="vertical", command=self.canvas.yview)
        self.scrollbar_h = ttk.Scrollbar(cf, orient="horizontal", command=self.canvas.xview)
        
        self.canvas.configure(yscrollcommand=self.scrollbar_v.set, xscrollcommand=self.scrollbar_h.set)
        
        # --- Packing order is important ---
        self.scrollbar_v.pack(side="right", fill="y")
        self.scrollbar_h.pack(side="bottom", fill="x")
        self.canvas.pack(side="left", fill="both", expand=True)

        # --- Frame inside Canvas ---
        self.scrollable_frame = ttk.Frame(self.canvas)
        self.canvas.create_window((0, 0), window=self.scrollable_frame, anchor="nw")

        self.scrollable_frame.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        
        # --- Mousewheel bindings ---
        self.bind_all("<MouseWheel>", lambda e: self.canvas.yview_scroll(-1 if e.delta > 0 else 1, "units"))
        self.bind_all("<Shift-MouseWheel>", lambda e: self.canvas.xview_scroll(-1 if e.delta > 0 else 1, "units"))

    def import_from_excel(self):
        fp = filedialog.askopenfilename(title="Pilih File Excel", filetypes=(("Excel Files", "*.xlsx"),))
        if not fp: return
        self.reset_ui(clear_log=False)
        try:
            wb = openpyxl.load_workbook(fp, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if len(row) < 4: continue
                name, ip, domain = (str(row[1] or "")).strip(), (str(row[2] or "")).strip(), (str(row[3] or "")).strip()
                if ip and not any(b.ip == ip for b in self.status_boxes):
                    self.add_status_box(ip, domain, name)
            self.rearrange_grid()
        except Exception as e:
            messagebox.showerror("Error Membaca File", f"Gagal mengimpor data:\n{e}")

    def add_manual_entry(self):
        d = ManualInputDialog(self)
        self.wait_window(d)
        if d.result:
            data = d.result
            ip = data['ip']
            if any(b.ip == ip for b in self.status_boxes):
                messagebox.showwarning("Duplikat", f"IP {ip} sudah ada.", parent=self)
                return
            self.add_status_box(ip, data['domain'], data['name'])
            self.rearrange_grid()

    def add_status_box(self, ip, domain, name):
        box = ClientStatusBox(self.scrollable_frame, ip, domain, name, self.log_activity)
        self.status_boxes.append(box)

    def rearrange_grid(self):
        try:
            cols = int(self.columns_var.get())
            if cols < 1: cols = 1
        except ValueError:
            cols = 5
        
        for i, b in enumerate(self.status_boxes):
            r, c = divmod(i, cols)
            b.grid(row=r, column=c, padx=5, pady=5, sticky="nsew")
        
        for c in range(cols):
            self.scrollable_frame.grid_columnconfigure(c, weight=1)
        self.update_control_states()

    def start_all(self):
        for b in self.status_boxes: b.start_check()

    def stop_all(self):
        for b in self.status_boxes: b.stop_check()

    def reset_ui(self, clear_log=True):
        self.stop_all()
        for b in self.status_boxes: b.destroy()
        self.status_boxes.clear()
        if clear_log:
            self.log_area.config(state='normal'); self.log_area.delete(1.0, tk.END); self.log_area.config(state='disabled')
        self.rearrange_grid()

    def update_control_states(self):
        state = "normal" if self.status_boxes else "disabled"
        self.start_all_button.config(state=state)
        self.stop_all_button.config(state=state)

    def export_results(self):
        if not self.status_boxes:
            messagebox.showwarning("Tidak Ada Data", "Tidak ada data siswa untuk diekspor.")
            return
        
        fp = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Workbook", "*.xlsx")], title="Simpan Laporan Hasil")
        if not fp: return
        
        try:
            wb = openpyxl.Workbook()
            sheet = wb.active
            sheet.title = "Hasil Penilaian"
            
            headers = ["No", "Nama Siswa", "Alamat IP", "Domain", "Nilai Akhir"] + list(self.status_boxes[0].check_items.values())
            sheet.append(headers)
            
            for i, b in enumerate(self.status_boxes, 1):
                status_texts = []
                for k in b.check_items:
                    v = b.results.get(k)
                    if v is True: status_texts.append("Berhasil")
                    elif v is False: status_texts.append("Gagal")
                    else: status_texts.append("-")
                
                row = [i, b.name, b.ip, b.domain_entry.get(), b.score] + status_texts
                sheet.append(row)
                
            wb.save(fp)
            messagebox.showinfo("Sukses", f"Hasil berhasil diekspor ke:\n{fp}")
        except Exception as e:
            messagebox.showerror("Error Ekspor", f"Gagal menyimpan file Excel:\n{e}")

    def log_activity(self, msg):
        ts = time.strftime("%Y-%m-%d %H:%M:%S")
        log = f"[{ts}] {msg}\n"
        self.log_area.config(state='normal')
        self.log_area.insert(tk.END, log)
        self.log_area.config(state='disabled')
        self.log_area.see(tk.END)

    def toggle_countdown(self):
        if self.timer_job_id:
            self.stop_countdown()
        else:
            try:
                m = int(self.timer_var.get())
                if m <= 0: return
                self.timer_seconds_remaining = m * 60
                self.timer_button.config(text="Stop Timer")
                self.timer_spinbox.config(state="disabled")
                self.update_countdown()
            except ValueError:
                self.timer_display_label.config(text="Error!")

    def update_countdown(self):
        if self.timer_seconds_remaining > 0:
            m, s = divmod(self.timer_seconds_remaining, 60)
            self.timer_display_label.config(text=f"{m:02d}:{s:02d}")
            self.timer_seconds_remaining -= 1
            self.timer_job_id = self.after(1000, self.update_countdown)
        else:
            self.timer_display_label.config(text="Selesai!")
            self.stop_countdown(triggered_by_finish=True)

    def stop_countdown(self, triggered_by_finish=False):
        if self.timer_job_id:
            self.after_cancel(self.timer_job_id)
            self.timer_job_id = None
        
        if triggered_by_finish:
            self.log_activity("TIMER SELESAI: Sesi monitoring dihentikan otomatis.")
            self.reset_ui()
        
        self.timer_button.config(text="Start Timer")
        self.timer_spinbox.config(state="normal")
        if not triggered_by_finish:
            self.timer_display_label.config(text="00:00")
            self.timer_seconds_remaining = 0

    def on_closing(self):
        self.stop_all()
        self.destroy()

if __name__ == "__main__":
    app = StatusMonitorApp()
    app.protocol("WM_DELETE_WINDOW", app.on_closing)
    app.mainloop()

